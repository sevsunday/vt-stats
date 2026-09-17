#!/usr/bin/env python3
"""
VT Stats -- F9bomber external-ledger importer (one-shot).

Reads F9bomber's hand-kept match spreadsheet (f9stats/*.xlsx), applies the
locked eligibility funnel, resolves identities against the project's
Steam64 registry + the committed alias table, pairs dual-recorded games
against our own telemetry corpus (ours supersedes), and writes two
committed artifacts:

  data/external/f9_ledger.json     normalized duels + overlap stamps +
                                   provenance (consumed by
                                   scripts/process_stats.py ->
                                   scripts/elo_commander.py and by the
                                   adjudication jogger)
  data/external/f9_community.json  display-only rollups (thug team
                                   records, commander records, faction
                                   picks, map counts) consumed 404-safe
                                   by js/player.js, js/app.js (Meta tab)
                                   and js/maps.js

THIS SCRIPT IS NOT PART OF THE PIPELINE. It runs once per ledger drop
(`python scripts/import_f9_ledger.py`), requires openpyxl (import-time
dependency only), and the pipeline reads the committed JSON forever
after. Decision memo: critique/decisions/f9-external-duels.md. Docs:
docs/DATA_DICTIONARY.md "External Community Ledger (F9bomber)".

Locked eligibility funnel (2026-09-13 operator sign-off; order matters
and the counters form a strict partition of the sheet's data rows):

  1. both Team One and Team Two lists present            (x_no_rosters)
  2. commanders parse as exactly "A vs B"                (x_bad_cmdr_parse)
     - names de-duped within each side; a commander accidentally listed
       in their own thug column is dropped from that list (notes only)
  3. even teams, thug count in {2, 3, 4} per side        (x_uneven / x_size)
     -> 3v3 / 4v4 / 5v5 lobbies only
  4. no participant doubles as a straggler               (x_straggler)
  5. winner is one of the two commanders                 (x_winner)
  6. duration present and >= MIN_DURATION_SEC            (x_nodur / x_short)
  7. exact-duplicate row dedup                           (x_exact_dup)
  8. overlap pairing vs our corpus -- ours supersedes    (g7_overlap)
     (paired rows become overlaps[] jogger entries regardless of
     excluded names; date +-1 day for the US-local vs UTC skew,
     same stripped map title, commander pair by Steam64 set when both
     resolve else keyname set, greedy one-to-one by participant-name
     Jaccard with floor OVERLAP_JACCARD_MIN)
  9. excluded-name filter on the remaining rows          (x_excluded_name)

Identity: exact norm() match against data/steamid_to_name.txt +
elo_current.json ratings names, with data/external/f9_name_map.json as
the override table. NEVER substring-matched; ambiguous (multi-Steam64)
names stay unresolved (steam64 null -> name-keyed ladder row).
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = PROJECT_ROOT / "f9stats" / "f9stats-20260913.xlsx"
EXTERNAL_DIR = PROJECT_ROOT / "data" / "external"
NAME_MAP_PATH = EXTERNAL_DIR / "f9_name_map.json"
MAP_ALIASES_PATH = EXTERNAL_DIR / "f9_map_aliases.json"
LEDGER_PATH = EXTERNAL_DIR / "f9_ledger.json"
COMMUNITY_PATH = EXTERNAL_DIR / "f9_community.json"
STEAMID_PATH = PROJECT_ROOT / "data" / "steamid_to_name.txt"
ELO_CURRENT_PATH = PROJECT_ROOT / "data" / "processed" / "elo_current.json"
MANIFEST_PATH = PROJECT_ROOT / "data" / "processed" / "matches.json"
REGISTRY_PATH = PROJECT_ROOT / "data" / "map-registry.json"

PROVIDER_NAME = "F9bomber"
PROVIDER_URL = "https://f9bomber.com"

# Locked gate constants (mirror scripts/elo.py's ELO_MIN_DURATION_SEC).
MIN_DURATION_SEC = 240
THUGS_PER_SIDE = (2, 3, 4)          # 3v3 / 4v4 / 5v5 with commanders
OVERLAP_DAY_TOLERANCE = 1           # F9 dates are US-local; ours are UTC
OVERLAP_JACCARD_MIN = 0.4

LEDGER_SCHEMA_VERSION = 1
COMMUNITY_SCHEMA_VERSION = 1

# Faction spellings seen on the sheet -> canonical name + project code.
FACTION_CANON = {
    "i.s.d.f": ("ISDF", "i"),
    "isdf": ("ISDF", "i"),
    "hadean": ("Hadean", "e"),
    "scion": ("Scion", "f"),
}


# ---------------------------------------------------------------------------
# Parsing helpers (validated against the sheet in the import investigation)
# ---------------------------------------------------------------------------

def norm(name: str | None) -> str:
    """Lowercase-alphanumeric name normalization. The ONLY comparison
    form used anywhere in this importer -- never substring-match."""
    return re.sub(r"[^a-z0-9]", "", (name or "").lower())


def parse_list(raw) -> list[str]:
    """Comma list, tolerating the sheet's `[a, b]` bracket style."""
    if raw is None:
        return []
    s = str(raw).strip()
    if not s:
        return []
    if s.startswith("[") and s.endswith("]"):
        s = s[1:-1]
    return [b.strip(" '\"") for b in re.split(r",\s*", s) if b.strip(" '\"")]


def parse_date(raw) -> datetime | None:
    """Sheet dates are `M.D.YY` (US-local calendar day)."""
    m = re.match(r"^(\d{1,2})\.(\d{1,2})\.(\d{2})$", str(raw).strip())
    if not m:
        return None
    return datetime(2000 + int(m.group(3)), int(m.group(1)), int(m.group(2)))


def parse_duration(raw) -> int | None:
    """`MM:SS` or `H:MM:SS` -> seconds."""
    if raw in (None, ""):
        return None
    m = re.match(r"^(\d+):(\d+)(?::(\d+))?$", str(raw).strip())
    if not m:
        return None
    a, b, c = int(m.group(1)), int(m.group(2)), int(m.group(3) or 0)
    return a * 3600 + b * 60 + c if m.group(3) is not None else a * 60 + b


def parse_commanders(raw) -> tuple[str | None, str | None]:
    parts = re.split(r"\s+vs\.?\s+", str(raw).strip(), flags=re.I)
    if len(parts) != 2:
        return None, None
    return parts[0].strip(), parts[1].strip()


def strip_map_title(name: str | None) -> str:
    """Map-title stem for joining F9 titles against registry titles and
    manifest display names (same rule the import investigation used)."""
    s = (name or "").lower().strip()
    s = re.sub(r"\.bzn$", "", s)
    s = re.sub(r"vsr$", "", s)
    s = re.sub(r"^(vsr:\s*|st:\s*|tvd:\s*)+", "", s)
    return re.sub(r"[^a-z0-9]", "", s)


def parse_faction(raw: str | None) -> tuple[str | None, str | None]:
    """-> (canonical_name, code) or (None, None)."""
    got = FACTION_CANON.get((raw or "").lower().strip())
    return got if got else (None, None)


# ---------------------------------------------------------------------------
# Identity resolution
# ---------------------------------------------------------------------------

class Resolver:
    """norm-name -> Steam64 resolution: alias table first, then a UNIQUE
    exact hit across steamid_to_name.txt + elo_current ratings names."""

    def __init__(self, name_map: dict[str, str], by_name: dict[str, set[str]],
                 canonical_by_s64: dict[str, str]):
        self.name_map = name_map
        self.by_name = by_name
        self.canonical_by_s64 = canonical_by_s64
        # Derived alias view for keyname(): normed F9 spelling -> normed
        # canonical spelling (so 'Blue Banana' and 'blue' collapse in
        # Jaccard sets / record merges even before Steam64s enter).
        self.name_aliases = {}
        for f9_norm, s64 in name_map.items():
            canon = canonical_by_s64.get(s64)
            if canon:
                self.name_aliases[f9_norm] = norm(canon)

    def keyname(self, name: str | None) -> str:
        k = norm(name)
        return self.name_aliases.get(k, k)

    def resolve(self, name: str | None) -> str | None:
        k = norm(name)
        if k in self.name_map:
            return self.name_map[k]
        hits = self.by_name.get(k, set())
        return next(iter(hits)) if len(hits) == 1 else None

    def display(self, name: str, s64: str | None) -> str:
        """Canonical display name when resolved, F9 spelling otherwise."""
        if s64 and self.canonical_by_s64.get(s64):
            return self.canonical_by_s64[s64]
        return name


def build_resolver(name_map: dict[str, str]) -> Resolver:
    by_name: dict[str, set[str]] = defaultdict(set)
    canonical_by_s64: dict[str, str] = {}
    for line in STEAMID_PATH.read_text(encoding="utf-8").splitlines():
        if "=" not in line:
            continue
        s64, name = line.split("=", 1)
        s64, name = s64.strip(), name.strip()
        if not (s64.isdigit() and name):
            continue
        by_name[norm(name)].add(s64)
        canonical_by_s64[s64] = name  # last wins (load_known_players parity)
    try:
        elo = json.loads(ELO_CURRENT_PATH.read_text(encoding="utf-8"))
        for r in elo.get("ratings", []):
            if r.get("steam64"):
                by_name[norm(r.get("name"))].add(str(r["steam64"]))
                canonical_by_s64.setdefault(str(r["steam64"]), r.get("name") or "")
    except (OSError, json.JSONDecodeError):
        print("WARN: elo_current.json unavailable; resolving from "
              "steamid_to_name.txt only")
    return Resolver(name_map, dict(by_name), canonical_by_s64)


# ---------------------------------------------------------------------------
# Import
# ---------------------------------------------------------------------------

def fail(msg: str) -> None:
    print(f"FATAL: {msg}")
    sys.exit(1)


def load_inputs():
    name_cfg = json.loads(NAME_MAP_PATH.read_text(encoding="utf-8"))
    map_cfg = json.loads(MAP_ALIASES_PATH.read_text(encoding="utf-8"))
    name_map = {norm(k): str(v) for k, v in (name_cfg.get("map") or {}).items()}
    exclude = {norm(x) for x in (name_cfg.get("exclude_names") or [])}
    map_aliases = {norm(k): str(v) for k, v in (map_cfg.get("map") or {}).items()}

    # Fail-loud: every alias Steam64 must exist in the registry file.
    steam_text = STEAMID_PATH.read_text(encoding="utf-8")
    known_s64 = {line.split("=", 1)[0].strip()
                 for line in steam_text.splitlines() if "=" in line}
    for f9_norm, s64 in name_map.items():
        if s64 not in known_s64:
            fail(f"f9_name_map entry {f9_norm!r} -> {s64} is not in "
                 f"{STEAMID_PATH.name}; add the canonical line first")

    registry = json.loads(REGISTRY_PATH.read_text(encoding="utf-8"))
    reg_maps = registry.get("maps") or registry
    if not isinstance(reg_maps, dict):
        fail("map-registry.json has an unexpected shape")
    for f9_stem, key in map_aliases.items():
        if key not in reg_maps:
            fail(f"f9_map_aliases entry {f9_stem!r} -> {key!r} is not a "
                 f"map-registry.json key")

    # Title-stem -> registry key index (alias table wins; first hit on a
    # stem collision, deterministic via sorted keys).
    title_index: dict[str, str] = {}
    for key in sorted(reg_maps):
        entry = reg_maps[key]
        if not isinstance(entry, dict):
            continue
        stem = strip_map_title(entry.get("title") or "")
        if stem and stem not in title_index:
            title_index[stem] = key
        kstem = strip_map_title(key)
        if kstem and kstem not in title_index:
            title_index[kstem] = key

    manifest = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
    return name_map, exclude, map_aliases, title_index, manifest


def read_rows(xlsx_path: Path) -> list[dict]:
    try:
        import openpyxl
    except ImportError:
        fail("openpyxl is required for the one-shot import "
             "(pip install openpyxl); the pipeline itself never needs it")
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, 12)]
    expected = ["Date", "Map", "Commanders", "Factions", "Winner",
                "Winning Faction", "Time", "Team One", "Team Two",
                "Straggler 1", "Straggler 2"]
    if headers != expected:
        fail(f"unexpected sheet headers {headers!r}; expected {expected!r}")
    rows = []
    for r in range(2, ws.max_row + 1):
        rec = {headers[c - 1]: ws.cell(r, c).value for c in range(1, 12)}
        rec["_row"] = r
        rows.append(rec)
    return rows


def resolve_map_key(map_title: str, map_aliases: dict, title_index: dict):
    """Registry key for an F9 title: alias table first, then the registry
    title-stem index. None when the map is unknown to us. Module-level so
    scripts/apply_f9_adjudications.py shares the exact join rule."""
    return (map_aliases.get(norm(map_title))
            or title_index.get(strip_map_title(map_title)))


def build_eligible_rows(rows: list[dict], resolver: Resolver
                        ) -> tuple[list[dict], Counter, Counter]:
    """Gates 1-7 (everything BEFORE overlap pairing): returns
    `(stage, funnel, notes)` where `stage` is the pre-overlap eligible row
    list (`{row, date, map_title, c1, c2, winner, duration_sec, t1, t2,
    factions_raw}`). Shared by run_import() and the one-shot
    scripts/apply_f9_adjudications.py (which re-derives the eligible set
    for its pairing-uniqueness gates)."""
    funnel: Counter = Counter()
    funnel["rows_total"] = len(rows)
    notes: Counter = Counter()

    stage: list[dict] = []
    seen_keys: set = set()
    for rec in rows:
        t1 = parse_list(rec["Team One"])
        t2 = parse_list(rec["Team Two"])
        if not (t1 and t2):
            funnel["x_no_rosters"] += 1
            continue
        c1, c2 = parse_commanders(rec["Commanders"])
        if not (c1 and c2):
            funnel["x_bad_cmdr_parse"] += 1
            continue
        u1 = list(dict.fromkeys(t1))
        u2 = list(dict.fromkeys(t2))
        if len(u1) != len(t1) or len(u2) != len(t2):
            notes["note_dup_name_in_team"] += 1
        if norm(c1) in {norm(x) for x in u1} or norm(c2) in {norm(x) for x in u2}:
            notes["note_cmdr_in_thugs"] += 1
            u1 = [x for x in u1 if norm(x) != norm(c1)]
            u2 = [x for x in u2 if norm(x) != norm(c2)]
        if len(u1) != len(u2):
            funnel["x_uneven"] += 1
            continue
        if len(u1) not in THUGS_PER_SIDE:
            funnel["x_size"] += 1
            continue
        stragglers = {norm(x) for x in
                      parse_list(rec["Straggler 1"]) + parse_list(rec["Straggler 2"])}
        participants_norm = {norm(x) for x in u1 + u2} | {norm(c1), norm(c2)}
        if participants_norm & stragglers:
            funnel["x_straggler"] += 1
            continue
        winner = str(rec["Winner"]).strip() if rec["Winner"] else ""
        if norm(winner) not in (norm(c1), norm(c2)):
            funnel["x_winner"] += 1
            continue
        dur = parse_duration(rec["Time"])
        if dur is None:
            funnel["x_nodur"] += 1
            continue
        if dur < MIN_DURATION_SEC:
            funnel["x_short"] += 1
            continue
        date = parse_date(rec["Date"])
        if date is None:
            # Never observed on the real sheet; counted defensively so the
            # partition can't silently leak.
            funnel["x_bad_date"] += 1
            continue
        dedup_key = (date, strip_map_title(rec["Map"]),
                     resolver.keyname(c1), resolver.keyname(c2),
                     tuple(sorted(resolver.keyname(x) for x in u1)),
                     tuple(sorted(resolver.keyname(x) for x in u2)),
                     dur, resolver.keyname(winner))
        if dedup_key in seen_keys:
            funnel["x_exact_dup"] += 1
            continue
        seen_keys.add(dedup_key)
        facs = parse_list(rec["Factions"])
        stage.append({
            "row": rec["_row"],
            "date": date,
            "map_title": (rec["Map"] or "").strip(),
            "c1": c1, "c2": c2,
            "winner": winner,
            "duration_sec": dur,
            "t1": u1, "t2": u2,
            "factions_raw": facs if len(facs) == 2 else [None, None],
        })
    funnel["g6_pre_overlap"] = len(stage)
    return stage, funnel, notes


def run_import(xlsx_path: Path) -> tuple[dict, dict]:
    name_map, exclude, map_aliases, title_index, manifest = load_inputs()
    resolver = build_resolver(name_map)
    rows = read_rows(xlsx_path)

    # ---- gates 1-7 (pre-overlap) ---------------------------------------
    stage, funnel, notes = build_eligible_rows(rows, resolver)

    # ---- gate 8: overlap pairing vs our corpus (ours supersedes) --------
    ours = []
    for m in manifest:
        dt = datetime.fromisoformat(m["date"].replace("Z", "+00:00"))
        leaders = m.get("team_leaders") or {}
        map_file_stem = re.sub(r"\.bzn$", "", str(m.get("map") or "").lower())
        ours.append({
            "id": m["id"],
            "date": dt.date(),
            "map_stem": strip_map_title(m.get("name") or ""),
            "map_file_stem": map_file_stem,
            "leader_s64s": {str((leaders.get("1") or {}).get("s64") or ""),
                            str((leaders.get("2") or {}).get("s64") or "")} - {""},
            "leader_names": {resolver.keyname((leaders.get("1") or {}).get("name")),
                             resolver.keyname((leaders.get("2") or {}).get("name"))} - {""},
            "players": {resolver.keyname(x) for x in (m.get("players") or [])},
            "decided_by": m.get("winner_decided_by"),
            "leaders": leaders,
        })

    candidates = []
    for i, p in enumerate(stage):
        f9_s64s = {resolver.resolve(p["c1"]), resolver.resolve(p["c2"])} - {None}
        f9_names = {resolver.keyname(p["c1"]), resolver.keyname(p["c2"])}
        all_names = {resolver.keyname(x)
                     for x in [p["c1"], p["c2"], *p["t1"], *p["t2"]]}
        p_map_key = resolve_map_key(p["map_title"], map_aliases, title_index)
        for j, o in enumerate(ours):
            if abs((o["date"] - p["date"].date()).days) > OVERLAP_DAY_TOLERANCE:
                continue
            # Map condition: display-title stems match, OR the F9 title's
            # resolved registry key equals our raw map-file stem (catches
            # community names -- "Quarry" vs "Quarry 2", "Mort's
            # Wasteland" vs "Wasteland", "Alien Dunes" vs "Patton's
            # Proving Ground" -- proven dual-records in the first import).
            stem_match = o["map_stem"] == strip_map_title(p["map_title"])
            key_match = bool(p_map_key) and p_map_key == o["map_file_stem"]
            if not (stem_match or key_match):
                continue
            pair_ok = ((len(f9_s64s) == 2 and f9_s64s == o["leader_s64s"])
                       or f9_names == o["leader_names"])
            if not pair_ok:
                continue
            inter = len(all_names & o["players"])
            union = len(all_names | o["players"]) or 1
            candidates.append((inter / union, i, j))
    candidates.sort(key=lambda t: (-t[0], t[1], t[2]))

    paired_f9: dict[int, tuple[float, int]] = {}
    used_ours: set[int] = set()
    for jac, i, j in candidates:
        if i in paired_f9 or j in used_ours or jac < OVERLAP_JACCARD_MIN:
            continue
        paired_f9[i] = (jac, j)
        used_ours.add(j)
    funnel["g7_overlap"] = len(paired_f9)

    overlaps = []
    for i in sorted(paired_f9):
        jac, j = paired_f9[i]
        p, o = stage[i], ours[j]
        winner_s64 = resolver.resolve(p["winner"])
        winner_key = resolver.keyname(p["winner"])
        our_team = None
        for team_key in ("1", "2"):
            leader = o["leaders"].get(team_key) or {}
            if winner_s64 and str(leader.get("s64") or "") == winner_s64:
                our_team = int(team_key)
                break
            if resolver.keyname(leader.get("name")) == winner_key:
                our_team = int(team_key)
                break
        overlaps.append({
            "f9_row": p["row"],
            "match_id": o["id"],
            "jaccard": round(jac, 3),
            "f9_winner_name": p["winner"],
            "our_team": our_team,
            "our_decided_by_at_import": o["decided_by"],
        })

    # ---- gate 9: excluded-name filter (unpaired rows only) --------------
    duels = []
    excluded_by_name: Counter = Counter()
    for i, p in enumerate(stage):
        if i in paired_f9:
            continue
        hit = None
        for x in [p["c1"], p["c2"], *p["t1"], *p["t2"]]:
            if norm(x) in exclude:
                hit = norm(x)
                break
        if hit:
            funnel["x_excluded_name"] += 1
            excluded_by_name[hit] += 1
            continue
        duels.append(p)
    funnel["FINAL"] = len(duels)

    # Partition check: every sheet row lands in exactly one bucket.
    partition = (funnel["x_no_rosters"] + funnel["x_bad_cmdr_parse"]
                 + funnel["x_uneven"] + funnel["x_size"]
                 + funnel["x_straggler"] + funnel["x_winner"]
                 + funnel["x_nodur"] + funnel["x_short"]
                 + funnel["x_bad_date"] + funnel["x_exact_dup"]
                 + funnel["g7_overlap"] + funnel["x_excluded_name"]
                 + funnel["FINAL"])
    if partition != funnel["rows_total"]:
        fail(f"funnel partition {partition} != rows_total {funnel['rows_total']}")

    # ---- ledger duels --------------------------------------------------
    def side_payload(names: list[str]) -> list[dict]:
        out = []
        for nm in names:
            s64 = resolver.resolve(nm)
            out.append({"name": nm, "steam64": s64})
        return out

    ledger_duels = []
    for p in sorted(duels, key=lambda d: (d["date"], d["row"])):
        winner_side = 1 if norm(p["winner"]) == norm(p["c1"]) else 2
        f1_name, _f1_code = parse_faction(p["factions_raw"][0])
        f2_name, _f2_code = parse_faction(p["factions_raw"][1])
        map_key = resolve_map_key(p["map_title"], map_aliases, title_index)
        n = len(p["t1"]) + 1
        ledger_duels.append({
            "row": p["row"],
            "date": p["date"].strftime("%Y-%m-%d"),
            "map_title": p["map_title"],
            "map_key": map_key,
            "size": f"{n}v{n}",
            "duration_sec": p["duration_sec"],
            "commanders": {
                "1": {"name": p["c1"], "steam64": resolver.resolve(p["c1"])},
                "2": {"name": p["c2"], "steam64": resolver.resolve(p["c2"])},
            },
            "thugs": {"1": side_payload(p["t1"]), "2": side_payload(p["t2"])},
            "winner_side": winner_side,
            "factions": {"1": f1_name, "2": f2_name},
        })

    # Per-duel invariants (fail loud; the _investigation gate re-checks).
    for d in ledger_duels:
        if d["size"] not in ("3v3", "4v4", "5v5"):
            fail(f"duel r{d['row']} has size {d['size']}")
        if len(d["thugs"]["1"]) != len(d["thugs"]["2"]):
            fail(f"duel r{d['row']} is uneven")
        if d["duration_sec"] < MIN_DURATION_SEC:
            fail(f"duel r{d['row']} under {MIN_DURATION_SEC}s")
        if d["winner_side"] not in (1, 2):
            fail(f"duel r{d['row']} has no winner_side")

    sheet_bytes = xlsx_path.read_bytes()
    ledger = {
        "schema_version": LEDGER_SCHEMA_VERSION,
        "provenance": {
            "provider": PROVIDER_NAME,
            "url": PROVIDER_URL,
            "source_file": xlsx_path.relative_to(PROJECT_ROOT).as_posix(),
            "sheet_sha256": hashlib.sha256(sheet_bytes).hexdigest(),
            "imported_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
            "importer": "scripts/import_f9_ledger.py",
        },
        "funnel": {**{k: funnel[k] for k in sorted(funnel)}, **dict(notes)},
        "duels": ledger_duels,
        "overlaps": overlaps,
        "excluded": {"by_name": {k: excluded_by_name[k]
                                 for k in sorted(excluded_by_name)}},
    }

    # ---- community rollups ----------------------------------------------
    thug_rec: dict[str, dict] = {}
    cmdr_rec: dict[str, dict] = {}

    def bump(store: dict, key: str, name: str, s64: str | None, won: bool):
        row = store.setdefault(key, {"steam64": s64, "name": name,
                                     "team_wins": 0, "team_losses": 0})
        row["team_wins" if won else "team_losses"] += 1

    for d in ledger_duels:
        win = d["winner_side"]
        for side in (1, 2):
            won = side == win
            c = d["commanders"][str(side)]
            ckey = c["steam64"] or f"name:{resolver.keyname(c['name'])}"
            bump(cmdr_rec, ckey, resolver.display(c["name"], c["steam64"]),
                 c["steam64"], won)
            for t in d["thugs"][str(side)]:
                tkey = t["steam64"] or f"name:{resolver.keyname(t['name'])}"
                bump(thug_rec, tkey, resolver.display(t["name"], t["steam64"]),
                     t["steam64"], won)

    def record_rows(store: dict, win_field: str, loss_field: str) -> list[dict]:
        out = []
        for row in store.values():
            games = row["team_wins"] + row["team_losses"]
            out.append({
                "steam64": row["steam64"],
                "name": row["name"],
                win_field: row["team_wins"],
                loss_field: row["team_losses"],
                "games": games,
            })
        out.sort(key=lambda r: (-r["games"], r["name"].lower()))
        return out

    faction_stats = {code: {"picks": 0, "wins": 0} for code in ("i", "e", "f")}
    for d in ledger_duels:
        for side in (1, 2):
            fac_name = d["factions"][str(side)]
            _canon, code = parse_faction(fac_name)
            if not code:
                continue
            faction_stats[code]["picks"] += 1
            if side == d["winner_side"]:
                faction_stats[code]["wins"] += 1

    map_counts: dict[tuple, dict] = {}
    for d in ledger_duels:
        key = (d["map_key"], d["map_title"] if d["map_key"] is None else "")
        row = map_counts.setdefault(key, {"map_key": d["map_key"],
                                          "title": d["map_title"], "games": 0})
        row["games"] += 1
    maps_out = sorted(map_counts.values(),
                      key=lambda r: (-r["games"], r["title"].lower()))

    dates = [d["date"] for d in ledger_duels]
    community = {
        "schema_version": COMMUNITY_SCHEMA_VERSION,
        "provider": {"name": PROVIDER_NAME, "url": PROVIDER_URL},
        "generated_at": ledger["provenance"]["imported_at"],
        "duel_count": len(ledger_duels),
        "date_range": [min(dates), max(dates)] if dates else [None, None],
        "thug_records": record_rows(thug_rec, "team_wins", "team_losses"),
        "commander_records": record_rows(cmdr_rec, "wins", "losses"),
        "faction_stats": faction_stats,
        "maps": maps_out,
    }
    return ledger, community


def main() -> None:
    ap = argparse.ArgumentParser(description="One-shot F9bomber ledger import")
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX,
                    help="Path to the F9 spreadsheet (default: %(default)s)")
    args = ap.parse_args()
    if not args.xlsx.exists():
        fail(f"{args.xlsx} not found")

    ledger, community = run_import(args.xlsx)

    EXTERNAL_DIR.mkdir(parents=True, exist_ok=True)
    with open(LEDGER_PATH, "w", encoding="utf-8") as f:
        json.dump(ledger, f, indent=2, ensure_ascii=False)
        f.write("\n")
    with open(COMMUNITY_PATH, "w", encoding="utf-8") as f:
        json.dump(community, f, indent=2, ensure_ascii=False)
        f.write("\n")

    fn = ledger["funnel"]
    print("=== F9 ledger import ===")
    print(f"  sheet: {ledger['provenance']['source_file']} "
          f"(sha256 {ledger['provenance']['sheet_sha256'][:16]}...)")
    print("  funnel:")
    for k in ("rows_total", "x_no_rosters", "x_bad_cmdr_parse", "x_uneven",
              "x_size", "x_straggler", "x_winner", "x_nodur", "x_short",
              "x_bad_date", "x_exact_dup", "g6_pre_overlap", "g7_overlap",
              "x_excluded_name", "FINAL"):
        print(f"    {k:18s} {fn.get(k, 0)}")
    unresolved_c = sum(1 for d in ledger["duels"]
                       for s in ("1", "2")
                       if not d["commanders"][s]["steam64"])
    print(f"  duels: {len(ledger['duels'])}  "
          f"(commander slots unresolved: {unresolved_c})")
    print(f"  overlaps recorded for adjudication jogger: {len(ledger['overlaps'])}")
    print(f"  wrote {LEDGER_PATH.relative_to(PROJECT_ROOT)}")
    print(f"  wrote {COMMUNITY_PATH.relative_to(PROJECT_ROOT)}")


if __name__ == "__main__":
    main()
